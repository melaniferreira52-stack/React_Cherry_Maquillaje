import io
from datetime import date, datetime, time

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.orm import Session, joinedload

from ..database import get_db
from ..dependencies import require_roles
from ..models import DetalleVenta, Usuario, Venta

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])

NOMBRE_PROYECTO = "Cherry Beauty — Maquillaje y Belleza"


def _ventas_del_dia(db: Session, fecha: date) -> list[Venta]:
    ventas = (
        db.query(Venta)
        .options(joinedload(Venta.detalles))
        .filter(
            Venta.fecha_hora >= datetime.combine(fecha, time.min),
            Venta.fecha_hora <= datetime.combine(fecha, time.max),
            Venta.estado == "registrada",
        )
        .order_by(Venta.fecha_hora.asc())
        .all()
    )
    return ventas


def _datos_reporte(ventas: list[Venta]) -> dict:
    """Estructura común usada por JSON, PDF y Excel (req. 4)."""
    filas = []
    for v in ventas:
        detalle = "; ".join(
            f"{d.nombre_item} x{d.cantidad}" for d in v.detalles
        ) if v.detalles else "—"
        filas.append(
            {
                "numero_venta": v.id,
                "fecha_hora": v.fecha_hora.strftime("%d/%m/%Y %H:%M") if v.fecha_hora else "",
                "cliente": f"{v.cliente_nombre} {v.cliente_apellido or ''}".strip(),
                "detalle": detalle,
                "total": float(v.total),
                "estado": v.estado,
            }
        )
    total_ventas = sum(f["total"] for f in filas)
    return {
        "proyecto": NOMBRE_PROYECTO,
        "filas": filas,
        "total_ventas": total_ventas,
        "cantidad_ventas": len(filas),
    }


@router.get("/ventas-diarias")
def reporte_ventas_diarias(
    fecha: date = Query(...),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    """Reporte diario de ventas en JSON (req. 4)."""
    ventas = _ventas_del_dia(db, fecha)
    datos = _datos_reporte(ventas)
    return {
        "proyecto": datos["proyecto"],
        "fecha_reporte": fecha.isoformat(),
        "ventas": datos["filas"],
        "totales": {
            "cantidad_ventas": datos["cantidad_ventas"],
            "total_ventas": datos["total_ventas"],
        },
        "generado_en": datetime.now().isoformat(),
    }


@router.get("/ventas-diarias/pdf")
def reporte_ventas_diarias_pdf(
    fecha: date = Query(...),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    """Exportación del reporte diario en PDF (req. 5)."""
    ventas = _ventas_del_dia(db, fecha)
    datos = _datos_reporte(ventas)

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="reportlab no está instalado")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"Reporte diario de ventas {fecha}",
    )

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("Titulo", parent=estilos["Title"], fontSize=16, textColor=colors.HexColor("#9f1239"))
    subtitulo = ParagraphStyle("Sub", parent=estilos["Normal"], fontSize=9, textColor=colors.HexColor("#57534e"))
    normal = estilos["Normal"]
    celda = ParagraphStyle("Celda", parent=normal, fontSize=8)

    elementos = []
    elementos.append(Paragraph("Cherry Beauty", titulo))
    elementos.append(Paragraph("Reporte diario de ventas", subtitulo))
    elementos.append(Paragraph(f"Fecha del reporte: {fecha.strftime('%d/%m/%Y')}", subtitulo))
    elementos.append(Spacer(1, 5 * mm))

    encabezados = ["N° venta", "Fecha y hora", "Cliente", "Productos / servicios", "Total", "Estado"]
    filas = [encabezados]
    for f in datos["filas"]:
        filas.append([
            str(f["numero_venta"]),
            f["fecha_hora"],
            f["cliente"],
            Paragraph(f["detalle"], celda),
            f"${f['total']:,.0f}".replace(",", "."),
            f["estado"],
        ])
    tabla = Table(filas, colWidths=[20 * mm, 38 * mm, 50 * mm, 95 * mm, 30 * mm, 25 * mm])
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ffe4e6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#9f1239")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (4, 0), (4, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#fecdd3")),
            ]
        )
    )
    elementos.append(tabla)
    elementos.append(Spacer(1, 5 * mm))

    resumen = Table(
        [
            ["Total de ventas", str(datos["cantidad_ventas"])],
            ["Valor total del día", f"${datos['total_ventas']:,.0f}".replace(",", ".")],
        ],
        colWidths=[80 * mm, 50 * mm],
    )
    resumen.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, -1), (-1, -1), colors.HexColor("#9f1239")),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ]
        )
    )
    elementos.append(resumen)
    elementos.append(Spacer(1, 6 * mm))
    elementos.append(
        Paragraph(
            f"Reporte generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} — Cherry Beauty",
            subtitulo,
        )
    )

    doc.build(elementos)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="reporte_ventas_{fecha}.pdf"'},
    )


@router.get("/ventas-diarias/excel")
def reporte_ventas_diarias_excel(
    fecha: date = Query(...),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    """Exportación del reporte diario en Excel (.xlsx) (req. 6)."""
    ventas = _ventas_del_dia(db, fecha)
    datos = _datos_reporte(ventas)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no está instalado")

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Reporte diario"

    hoja.append(["Cherry Beauty", "", "", "", "", ""])
    hoja.append(["Reporte diario de ventas", "", "", "", "", ""])
    hoja.append([f"Fecha del reporte: {fecha.strftime('%d/%m/%Y')}", "", "", "", "", ""])
    hoja.append([])

    encabezados = ["N° venta", "Fecha y hora", "Cliente", "Productos / servicios", "Total", "Estado"]
    hoja.append(encabezados)
    for celda in hoja[5]:
        celda.font = Font(bold=True, color="9F1239")
        celda.fill = PatternFill("solid", fgColor="FFE4E6")
        celda.alignment = Alignment(horizontal="center")

    for f in datos["filas"]:
        hoja.append([
            f["numero_venta"],
            f["fecha_hora"],
            f["cliente"],
            f["detalle"],
            f["total"],
            f["estado"],
        ])

    hoja.append([])
    hoja.append(["Total de ventas", datos["cantidad_ventas"], "", "", "", ""])
    hoja.append(["Valor total del día", datos["total_ventas"], "", "", "", ""])
    hoja.append([])
    hoja.append([f"Reporte generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"])

    for fila in hoja.iter_rows(min_row=6, min_col=1, max_col=6):
        for celda in fila:
            celda.alignment = Alignment(vertical="top")

    hoja.column_dimensions["A"].width = 12
    hoja.column_dimensions["B"].width = 20
    hoja.column_dimensions["C"].width = 32
    hoja.column_dimensions["D"].width = 60
    hoja.column_dimensions["E"].width = 14
    hoja.column_dimensions["F"].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="reporte_ventas_{fecha}.xlsx"'},
    )